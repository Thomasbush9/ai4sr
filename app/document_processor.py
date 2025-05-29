from typing import List, Dict, Any
import os
from pathlib import Path
from langchain.text_splitter import RecursiveCharacterTextSplitter
from langchain_openai import AzureOpenAIEmbeddings
import fitz  # PyMuPDF
import docx
import json
import logging
import uuid
from sqlalchemy.orm import Session
from models import Document, DocumentChunk, User
from azure_search import AzureSearchManager

logger = logging.getLogger(__name__)

class DocumentProcessor:
    def __init__(self, embedding_config: Dict[str, str], azure_search_config: Dict[str, str], db_session: Session):
        """Initialize with embedding and Azure Search configuration."""
        self.embedding_config = embedding_config
        self.embeddings = AzureOpenAIEmbeddings(
            model=embedding_config["deployment_name"],
            azure_endpoint=embedding_config["api_base"],
            api_key=embedding_config["api_key"],
            api_version=embedding_config["api_version"]
        )
        self.text_splitter = RecursiveCharacterTextSplitter(
            chunk_size=1000,
            chunk_overlap=200,
            length_function=len
        )
        self.azure_search = AzureSearchManager(
            endpoint=azure_search_config["endpoint"],
            key=azure_search_config["key"]
        )
        self.db_session = db_session
        logger.info(f"Initialized DocumentProcessor with embedding model: {embedding_config['deployment_name']}")
        logger.info(f"Using embedding endpoint: {embedding_config['api_base']}")

    def extract_text(self, file_path: str) -> str:
        """Extract text from various document formats."""
        file_extension = Path(file_path).suffix.lower()
        logger.info(f"Extracting text from {file_path} with extension {file_extension}")
        
        if file_extension == '.pdf':
            return self._extract_from_pdf(file_path)
        elif file_extension == '.docx':
            return self._extract_from_docx(file_path)
        elif file_extension == '.txt':
            return self._extract_from_txt(file_path)
        else:
            raise ValueError(f"Unsupported file format: {file_extension}")

    def _extract_from_pdf(self, file_path: str) -> str:
        """Extract text from PDF file."""
        doc = fitz.open(file_path)
        text = ""
        for page in doc:
            text += page.get_text()
        logger.info(f"Extracted {len(text)} characters from PDF")
        return text

    def _extract_from_docx(self, file_path: str) -> str:
        """Extract text from DOCX file."""
        doc = docx.Document(file_path)
        text = "\n".join([paragraph.text for paragraph in doc.paragraphs])
        logger.info(f"Extracted {len(text)} characters from DOCX")
        return text

    def _extract_from_txt(self, file_path: str) -> str:
        """Extract text from TXT file."""
        with open(file_path, 'r', encoding='utf-8') as f:
            text = f.read()
        logger.info(f"Extracted {len(text)} characters from TXT")
        return text

    def create_chunks(self, text: str) -> List[str]:
        """Split text into chunks."""
        chunks = self.text_splitter.split_text(text)
        logger.info(f"Created {len(chunks)} chunks from text")
        return chunks

    def generate_embeddings(self, chunks: List[str], document_id: str, index_name: str) -> None:
        """Generate embeddings and store them in Azure Cognitive Search."""
        logger.info(f"Generating embeddings for {len(chunks)} chunks")
        try:
            # Generate embeddings for each chunk
            documents = []
            for i, chunk in enumerate(chunks):
                # Generate embedding
                embedding = self.embeddings.embed_query(chunk)
                
                # Create document for Azure Search
                doc = {
                    "id": str(uuid.uuid4()),
                    "content": chunk,
                    "document_id": document_id,
                    "chunk_index": i,
                    "embedding": embedding
                }
                documents.append(doc)
            
            # Upload documents to Azure Search
            self.azure_search.upload_documents(index_name, documents)
            logger.info(f"Uploaded {len(documents)} documents to Azure Search index: {index_name}")
            
        except Exception as e:
            logger.error(f"Error generating embeddings: {str(e)}")
            raise

    def process_document(self, file_path: str, user_id: int) -> Dict[str, Any]:
        """Process a document and store its embeddings."""
        logger.info(f"Processing document: {file_path}")
        
        try:
            # Extract text
            text = self.extract_text(file_path)
            
            # Create chunks
            chunks = self.create_chunks(text)
            
            # Create document record in database
            document = Document(
                user_id=user_id,
                filename=Path(file_path).name,
                file_path=str(file_path),
                file_type=Path(file_path).suffix.lower(),
                num_chunks=len(chunks),
                azure_search_index=f"doc_{user_id}_{Path(file_path).stem}"
            )
            self.db_session.add(document)
            self.db_session.commit()
            
            # Create Azure Search index
            self.azure_search.create_index(document.azure_search_index)
            
            # Generate and store embeddings
            self.generate_embeddings(chunks, str(document.id), document.azure_search_index)
            
            # Store chunks in database
            for i, chunk in enumerate(chunks):
                chunk_record = DocumentChunk(
                    document_id=document.id,
                    chunk_index=i,
                    content=chunk,
                    azure_search_id=str(uuid.uuid4())
                )
                self.db_session.add(chunk_record)
            self.db_session.commit()
            
            # Return metadata
            metadata = {
                "document_id": document.id,
                "filename": document.filename,
                "num_chunks": len(chunks),
                "azure_search_index": document.azure_search_index
            }
            logger.info(f"Successfully processed document: {metadata['filename']}")
            return metadata
            
        except Exception as e:
            self.db_session.rollback()
            logger.error(f"Error processing document: {str(e)}")
            raise 